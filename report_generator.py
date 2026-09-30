import os, json, platform, io, sys, tempfile

from providers import resolve_provider

def _config_path():
    s = platform.system()
    if s == 'Windows':
        base = os.environ.get('APPDATA', os.path.expanduser('~'))
    elif s == 'Darwin':
        base = os.path.expanduser('~/Library/Application Support')
    else:
        base = os.path.expanduser('~/.config')
    return os.path.join(base, 'PsychAid', 'config.json')

def load_config():
    p = _config_path()
    if os.path.exists(p):
        with open(p) as f:
            raw = json.load(f)
        cfg = _normalise_config(raw)
        if cfg != raw:
            save_config(cfg)
        return cfg
    return _normalise_config({})

def save_config(data):
    p = _config_path()
    directory = os.path.dirname(p)
    os.makedirs(directory, mode=0o700, exist_ok=True)
    fd, temp_path = tempfile.mkstemp(dir=directory)
    try:
        with os.fdopen(fd, 'w') as f:
            json.dump(_normalise_config(data), f)
        try:
            os.chmod(temp_path, 0o600)
        except OSError:
            pass
        os.replace(temp_path, p)
    except Exception:
        try:
            os.unlink(temp_path)
        except OSError:
            pass
        raise

def _normalise_config(data):
    data = data if isinstance(data, dict) else {}
    keys = data.get('api_keys') if isinstance(data.get('api_keys'), dict) else {}
    keys = {
        name: value for name, value in keys.items()
        if isinstance(name, str) and isinstance(value, str)
    }

    legacy_key = data.get('api_key')
    if isinstance(legacy_key, str) and legacy_key.strip() and not keys.get('anthropic'):
        keys['anthropic'] = legacy_key.strip()

    provider = data.get('provider') if isinstance(data.get('provider'), str) else 'anthropic'
    model = data.get('model') if isinstance(data.get('model'), str) else ''
    auth_mode = data.get('auth_mode') if data.get('auth_mode') in ('api_key', 'subscription') else 'api_key'
    return {
        'provider': provider or 'anthropic',
        'model': model.strip(),
        'auth_mode': auth_mode,
        'api_keys': keys
    }

def extract_text(file_bytes, filename):
    ext = os.path.splitext(filename)[1].lower()
    if ext == '.pdf':
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        return '\n'.join(p.extract_text() or '' for p in reader.pages)
    elif ext == '.docx':
        import mammoth
        result = mammoth.extract_raw_text(io.BytesIO(file_bytes))
        return result.value
    else:
        return file_bytes.decode('utf-8', errors='replace')

SYSTEM_PROMPT = """You are a clinical report assistant for an Australian psychology practice. 
You receive assessment documents, intake forms, and score reports uploaded by a psychologist. 
Your task is to synthesise the information into a structured psychological report.

Structure the report with these sections:
1. Identifying Information
2. Reason for Referral
3. Background History (developmental, medical, educational, family/social)
4. Assessment Measures Administered
5. Behavioural Observations (leave placeholder: [To be completed by clinician])
6. Assessment Results (leave placeholder: [To be completed by clinician])
7. Summary and Diagnostic Impressions (leave placeholder: [To be completed by clinician])
8. Recommendations

For the Recommendations section, include relevant recommendations based on the profile presented.
Use professional clinical language appropriate for Australian psychological practice.
Do not fabricate information not present in the source documents."""

_SKILL_PROMPT = None

def load_skill_prompt():
    global _SKILL_PROMPT
    if _SKILL_PROMPT is not None:
        return _SKILL_PROMPT

    if getattr(sys, 'frozen', False):
        base = sys._MEIPASS
    else:
        base = os.path.dirname(os.path.abspath(__file__))
    path = os.path.join(base, 'skill', '4thought-report', 'SKILL.md')

    try:
        with open(path, encoding='utf-8') as f:
            text = f.read()
        lines = text.splitlines(keepends=True)
        delimiters = [i for i, line in enumerate(lines) if line.strip() == '---']
        if delimiters and delimiters[0] == 0 and len(delimiters) > 1:
            text = ''.join(lines[delimiters[1] + 1:])
        _SKILL_PROMPT = text or SYSTEM_PROMPT
    except (OSError, UnicodeError):
        _SKILL_PROMPT = SYSTEM_PROMPT
    return _SKILL_PROMPT

def generate_report(files, cfg):
    all_text = []
    for fname, fbytes in files:
        try:
            text = extract_text(fbytes, fname)
            all_text.append(f"=== {fname} ===\n{text}")
        except Exception as e:
            all_text.append(f"=== {fname} ===\n[Error extracting: {e}]")
    
    combined = '\n\n'.join(all_text)
    
    provider, api_key, model = resolve_provider(cfg)
    return provider['generate'](
        load_skill_prompt(),
        f"Please generate a psychological report based on the following documents:\n\n{combined}",
        api_key,
        model
    )

def report_to_docx(text, client_name=''):
    """
    Convert AI-generated report markdown to a styled Word document.
    Uses the 4Thought template (report_template.docx) which defines:
      - 'paragraph' style  → cover page items (PSYCHOLOGICAL ASSESSMENT REPORT, CLIENT DETAILS, field labels)
      - 'Normal' bold      → ALL CAPS section headings
      - 'Normal'           → body text
      - 'ListParagraph'    → bullet items
      - 'Table Grid'       → score tables
    """
    import re
    from docx import Document
    from docx.shared import Pt
    from docx.oxml import OxmlElement

    # ── Load the branded template ────────────────────────────────────────────
    if getattr(sys, 'frozen', False):
        base = sys._MEIPASS
    else:
        base = os.path.dirname(os.path.abspath(__file__))
    template_path = os.path.join(base, 'report_template.docx')

    try:
        doc = Document(template_path)
    except Exception:
        doc = Document()

    # ── Clear body content, preserve header/footer (they live in sectPr refs) ─
    body = doc.element.body
    W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    for child in list(body):
        if child.tag in (f'{{{W_NS}}}p', f'{{{W_NS}}}tbl'):
            body.remove(child)

    # ── Ensure the logo header appears on ALL pages including page 1 ──────────
    # Find the body-level sectPr and locate the "default" header rId (the one
    # with the logo). Then add a matching "first" headerReference with the same
    # rId and set titlePg so Word uses it on page 1.
    sectPr = body.find(f'{{{W_NS}}}sectPr')
    if sectPr is not None:
        # Find the default header rId
        default_rid = None
        for href in sectPr.findall(f'{{{W_NS}}}headerReference'):
            htype = href.get(f'{{{W_NS}}}type', '')
            if htype == 'default':
                default_rid = href.get(f'{{{R_NS}}}id', '')
                break
        if default_rid:
            # Add a "first" page header pointing to the same logo header
            first_exists = any(
                h.get(f'{{{W_NS}}}type') == 'first'
                for h in sectPr.findall(f'{{{W_NS}}}headerReference')
            )
            if not first_exists:
                first_href = OxmlElement('w:headerReference')
                first_href.set(f'{{{W_NS}}}type', 'first')
                first_href.set(f'{{{R_NS}}}id', default_rid)
                sectPr.insert(0, first_href)
            # Enable "different first page" so Word uses the first-page header
            if sectPr.find(f'{{{W_NS}}}titlePg') is None:
                titlePg = OxmlElement('w:titlePg')
                sectPr.append(titlePg)

    # ── Style lookup helper ───────────────────────────────────────────────────
    def _style(name):
        try:
            return doc.styles[name]
        except Exception:
            return doc.styles['Normal']

    # ── Inline bold parser ────────────────────────────────────────────────────
    def _add_runs(para, text, default_bold=False):
        parts = re.split(r'(\*\*[^*]+\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = para.add_run(part[2:-2])
                run.bold = True
            elif part:
                run = para.add_run(part)
                run.bold = default_bold

    # ── Paragraph builders ────────────────────────────────────────────────────
    def add_cover(text, bold=False):
        """'paragraph' custom style — cover page items, centred."""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        p = doc.add_paragraph(style=_style('paragraph'))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_runs(p, text, default_bold=bold)
        return p

    def add_hrule():
        """Insert a VML horizontal rule matching the template (gray line across page)."""
        from lxml import etree
        hr_xml = (
            '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
            ' xmlns:v="urn:schemas-microsoft-com:vml"'
            ' xmlns:o="urn:schemas-microsoft-com:office:office">'
            '<w:pPr>'
            '<w:pStyle w:val="paragraph"/>'
            '<w:spacing w:before="120" w:beforeAutospacing="0"'
            ' w:after="120" w:afterAutospacing="0"/>'
            '<w:jc w:val="center"/>'
            '</w:pPr>'
            '<w:r>'
            '<w:rPr><w:rFonts w:ascii="Cambria" w:hAnsi="Cambria"/>'
            '<w:noProof/><w:color w:val="000000"/></w:rPr>'
            '<w:pict>'
            '<v:rect style="width:451.3pt;height:.05pt" fillcolor="#a0a0a0" stroked="f"'
            ' o:hr="t" o:hrstd="t" o:hralign="center"/>'
            '</w:pict>'
            '</w:r>'
            '</w:p>'
        )
        placeholder = doc.add_paragraph()
        new_el = etree.fromstring(hr_xml)
        placeholder._element.getparent().replace(placeholder._element, new_el)

    def add_client_field(label, value):
        """Bold label + normal value, left-aligned within the cover block."""
        p = doc.add_paragraph(style=_style('paragraph'))
        r = p.add_run(label)
        r.bold = True
        if value:
            p.add_run(' ' + value)
        return p

    # ASSESSMENTS lives on the cover page (centred, no page break).
    # The first *content* heading (REASON FOR REFERRAL and onwards) gets a page break.
    _first_content_heading_done = [False]

    def add_heading(text, cover=False):
        """Normal style + bold — ALL CAPS section headings."""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        p = doc.add_paragraph(style=_style('Normal'))
        run = p.add_run(text)
        run.bold = True
        p.paragraph_format.space_after = Pt(2)
        if cover:
            # ASSESSMENTS: centred on the cover page, no page break
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
        elif not _first_content_heading_done[0]:
            # First real content heading → start on a new page
            p.paragraph_format.page_break_before = True
            _first_content_heading_done[0] = True
        else:
            p.paragraph_format.space_before = Pt(10)
        return p

    def add_body(text):
        """Normal style — body paragraphs."""
        p = doc.add_paragraph(style=_style('Normal'))
        _add_runs(p, text)
        p.paragraph_format.space_after = Pt(6)
        return p

    def add_bullet(text):
        """List Paragraph style — bullet items."""
        p = doc.add_paragraph(style=_style('List Paragraph'))
        _add_runs(p, text)
        return p

    def add_table(rows):
        if not rows:
            return
        cols = max(len(r) for r in rows)
        try:
            table = doc.add_table(rows=len(rows), cols=cols, style='Table Grid')
        except Exception:
            table = doc.add_table(rows=len(rows), cols=cols)
        for ri, row_data in enumerate(rows):
            for ci in range(cols):
                cell_text = row_data[ci].strip() if ci < len(row_data) else ''
                cell = table.cell(ri, ci)
                cell.text = ''
                run = cell.paragraphs[0].add_run(cell_text)
                run.bold = (ri == 0)

    # ── Cover page items ──────────────────────────────────────────────────────
    COVER_ITEMS = {'PSYCHOLOGICAL ASSESSMENT REPORT', 'CONFIDENTIAL', 'CLIENT DETAILS'}
    # ASSESSMENTS stays on the cover page (centred, no page break before it)
    COVER_HEADINGS = {'ASSESSMENTS'}

    CLIENT_FIELD_LABELS = (
        'Client name:', 'Date of birth:', 'Age at time of testing:',
        'Gender:', 'School:', 'Date of assessment:', 'Date of report:',
        'Clinician:', 'Location of assessment:',
    )

    def _is_client_field(text):
        bare = re.sub(r'^\*+|\*+$', '', text).strip()
        return any(bare.startswith(lbl) for lbl in CLIENT_FIELD_LABELS)

    def _parse_client_field(text):
        bare = re.sub(r'^\*+|\*+$', '', text).strip()
        for lbl in CLIENT_FIELD_LABELS:
            if bare.startswith(lbl):
                value = bare[len(lbl):].strip().lstrip('*').strip()
                return lbl, value
        return None, text

    def _is_all_caps_heading(text):
        bare = re.sub(r'\*+', '', text).strip()
        return (bare.isupper() and len(bare) > 3
                and not bare.startswith('-')
                and not re.match(r'^\d+[\.\)]\s', bare))

    # ── Parse and render lines ────────────────────────────────────────────────
    # State tracking for horizontal rules around CLIENT DETAILS block
    _after_client_details = [False]   # we've seen CLIENT DETAILS heading
    _in_client_fields = [False]       # currently inside the client field block
    _client_fields_closed = [False]   # closing hrule already inserted

    lines = text.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # Skip markdown horizontal rules
        if re.match(r'^[-*]{3,}$', stripped):
            i += 1
            continue

        # Markdown table
        if stripped.startswith('|') and stripped.endswith('|'):
            # Close client-field block if still open
            if _in_client_fields[0] and not _client_fields_closed[0]:
                add_hrule()
                _in_client_fields[0] = False
                _client_fields_closed[0] = True
            table_rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = lines[i].strip()
                if re.match(r'^\|[\s|:-]+\|$', row):
                    i += 1
                    continue
                cells = [c.strip() for c in row.strip('|').split('|')]
                table_rows.append(cells)
                i += 1
            add_table(table_rows)
            continue

        # Strip markdown heading prefix for content
        clean = re.sub(r'^#{1,3}\s*', '', stripped).strip()
        bare = re.sub(r'^\*+|\*+$', '', clean).strip()

        # Cover page: PSYCHOLOGICAL ASSESSMENT REPORT / CONFIDENTIAL / CLIENT DETAILS
        if bare in COVER_ITEMS:
            bold = bare != 'CONFIDENTIAL'
            add_cover(bare, bold=bold)
            if bare == 'CLIENT DETAILS':
                _after_client_details[0] = True
            i += 1
            continue

        # Client detail field lines — open hrule block on first field
        if _is_client_field(clean):
            if _after_client_details[0] and not _in_client_fields[0]:
                add_hrule()          # opening rule above first field
                _in_client_fields[0] = True
            label, value = _parse_client_field(clean)
            if label:
                add_client_field(label, value)
            else:
                add_cover(clean)
            i += 1
            continue

        # Any non-field content after client fields → close hrule block
        if _in_client_fields[0] and not _client_fields_closed[0]:
            add_hrule()              # closing rule below last field
            _in_client_fields[0] = False
            _client_fields_closed[0] = True

        # ALL CAPS section heading or markdown heading
        is_md_heading = bool(re.match(r'^#{1,3}\s', stripped))
        if is_md_heading or _is_all_caps_heading(clean):
            is_cover_heading = bare in COVER_HEADINGS
            add_heading(bare, cover=is_cover_heading)
            i += 1
            continue

        # Bullet item
        if re.match(r'^[-*•]\s+', stripped):
            add_bullet(re.sub(r'^[-*•]\s+', '', stripped))
            i += 1
            continue
        if re.match(r'^\d+[\.\)]\s+', stripped):
            add_bullet(re.sub(r'^\d+[\.\)]\s+', '', stripped))
            i += 1
            continue

        # Body paragraph
        add_body(clean)
        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# ── STEP mini report ──────────────────────────────────────────────────────────

STEP_SYSTEM_PROMPT = """You are generating a STEP Assessment Report — a brief, non-diagnostic screening summary.

The STEP report must contain EXACTLY the following, in order:

1. A title line: STEP Assessment Report

2. BACKGROUND INFORMATION
   One short paragraph (2–4 sentences) covering only: client's age, gender, who referred them, and the primary reason for referral. Nothing else.

3. SCREENER SUMMARIES
   For each screener/questionnaire administered, write ONE paragraph that:
   - Names the screener and who completed it (e.g. parent, teacher, self)
   - States the raw scores or totals as reported in the score report
   - Describes what domain was measured and what ratings were endorsed (factual only)
   - Does NOT interpret, diagnose, or draw clinical conclusions

4. CONSISTENCY AND INCONSISTENCY SUMMARY
   One paragraph noting where screeners agree with each other and where they differ. Factual only — no diagnostic conclusions, no clinical interpretation.

Rules you must follow:
- This document is a SCREENING TOOL, not a clinical report
- Do NOT include diagnoses or diagnostic impressions
- Do NOT interpret scores clinically (no "elevated", "clinically significant", "meets criteria")
- Do NOT include recommendations
- Do NOT include background history beyond age/gender/referral reason
- Do NOT include cognitive assessment results, IQ scores, or index scores
- Keep language plain, professional, and factual
- Output headings in ALL CAPS, body text as plain paragraphs
- The entire document should be no longer than one page
"""


def generate_step_report(doc_text, cfg):
    """Generate a STEP screening summary from the raw document text."""
    from providers import resolve_provider
    provider, api_key, model = resolve_provider(cfg)
    return provider['generate'](
        STEP_SYSTEM_PROMPT,
        f"Generate a STEP Assessment Report from the following uploaded documents:\n\n{doc_text}",
        api_key,
        model
    )


def step_to_docx(text, client_name=''):
    """Convert STEP report text to a clean Word document."""
    import re
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Pt(72)
        section.bottom_margin = Pt(72)
        section.left_margin = Pt(90)
        section.right_margin = Pt(90)

    lines = text.strip().split('\n')
    title_done = False

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        # Strip markdown heading prefix
        clean = re.sub(r'^#{1,3}\s*', '', stripped).strip()
        bare = re.sub(r'^\*+|\*+$', '', clean).strip()

        # Title line
        if not title_done and 'STEP' in bare.upper() and 'ASSESSMENT' in bare.upper():
            p = doc.add_paragraph()
            run = p.add_run(bare)
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x1a, 0x2d, 0x4f)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(14)
            title_done = True
            continue

        # ALL CAPS section heading
        bare_plain = re.sub(r'\*+', '', bare).strip()
        if bare_plain.isupper() and len(bare_plain) > 3:
            p = doc.add_paragraph()
            run = p.add_run(bare_plain)
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x1a, 0x2d, 0x4f)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            continue

        # Bullet
        if re.match(r'^[-*•]\s+', stripped):
            p = doc.add_paragraph(style='List Bullet')
            content = re.sub(r'^[-*•]\s+', '', stripped)
            p.add_run(content)
            p.paragraph_format.space_after = Pt(2)
            continue

        # Body paragraph
        p = doc.add_paragraph()
        p.add_run(bare)
        p.paragraph_format.space_after = Pt(6)

    if client_name:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        run = p.add_run(f'Client: {client_name}')
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x94, 0xa3, 0xb8)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
