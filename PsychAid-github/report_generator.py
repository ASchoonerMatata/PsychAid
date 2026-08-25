import os, json, platform, io

def _config_path():
    s = platform.system()
    if s == 'Windows':
        base = os.environ.get('APPDATA', os.path.expanduser('~'))
    elif s == 'Darwin':
        base = os.path.expanduser('~/Library/Application Support')
    else:
        base = os.path.expanduser('~/.config')
    return os.path.join(base, 'PsychAid', 'config.json')

def load_config():
    p = _config_path()
    if os.path.exists(p):
        with open(p) as f:
            return json.load(f)
    return {}

def save_config(data):
    p = _config_path()
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, 'w') as f:
        json.dump(data, f)

def extract_text(file_bytes, filename):
    ext = os.path.splitext(filename)[1].lower()
    if ext == '.pdf':
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        return '\n'.join(p.extract_text() or '' for p in reader.pages)
    elif ext in ('.docx', '.doc'):
        import mammoth
        result = mammoth.extract_raw_text(io.BytesIO(file_bytes))
        return result.value
    else:
        return file_bytes.decode('utf-8', errors='replace')

SYSTEM_PROMPT = """You are a clinical report assistant for an Australian psychology practice. 
You receive assessment documents, intake forms, and score reports uploaded by a psychologist. 
Your task is to synthesise the information into a structured psychological report.

Structure the report with these sections:
1. Identifying Information
2. Reason for Referral
3. Background History (developmental, medical, educational, family/social)
4. Assessment Measures Administered
5. Behavioural Observations (leave placeholder: [To be completed by clinician])
6. Assessment Results (leave placeholder: [To be completed by clinician])
7. Summary and Diagnostic Impressions (leave placeholder: [To be completed by clinician])
8. Recommendations

For the Recommendations section, include relevant recommendations based on the profile presented.
Use professional clinical language appropriate for Australian psychological practice.
Do not fabricate information not present in the source documents."""

def generate_report(files, api_key):
    import anthropic
    client = anthropic.Anthropic(api_key=api_key)
    
    all_text = []
    for fname, fbytes in files:
        try:
            text = extract_text(fbytes, fname)
            all_text.append(f"=== {fname} ===\n{text}")
        except Exception as e:
            all_text.append(f"=== {fname} ===\n[Error extracting: {e}]")
    
    combined = '\n\n'.join(all_text)
    
    response = client.messages.create(
        model='claude-opus-4-5',
        max_tokens=8000,
        system=SYSTEM_PROMPT,
        messages=[{
            'role': 'user',
            'content': f"Please generate a psychological report based on the following documents:\n\n{combined}"
        }]
    )
    return response.content[0].text

def report_to_docx(text, client_name=''):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    
    # Title
    title = doc.add_heading('PSYCHOLOGICAL ASSESSMENT REPORT', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    if client_name:
        sub = doc.add_paragraph(f'Prepared for: {client_name}')
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('')
    
    # Parse and add content
    for line in text.split('\n'):
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph('')
            continue
        # Detect headings (numbered sections or ALL CAPS)
        if (stripped[0].isdigit() and '. ' in stripped[:4]) or stripped.isupper():
            doc.add_heading(stripped, level=1)
        else:
            doc.add_paragraph(stripped)
    
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
